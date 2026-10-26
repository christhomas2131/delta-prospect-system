"""Generate portable company intelligence reports in PDF, Word, and CSV."""

import csv
import io
import json
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal
from html import escape
from typing import Any


def safe_filename(value: str | None) -> str:
    """Return a portable filename stem without trusting database text."""
    normalized = unicodedata.normalize("NFKD", value or "company").encode("ascii", "ignore").decode()
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", normalized).strip("._-")
    return cleaned or "company"


def display_value(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, indent=2, default=str, ensure_ascii=False)
    return str(value)


def _prize_summary(prospect: dict) -> str:
    pb = prospect.get("prize_breakdown") or {}
    if isinstance(pb, str):
        try:
            pb = json.loads(pb)
        except ValueError:
            pb = {}
    rng = pb.get("range") or {}
    if not rng:
        return display_value(prospect.get("size_of_prize"))
    fee = pb.get("fee_estimate") or {}
    parts = [f"{pb.get('currency', 'AUD')} {rng.get('low'):,} – {rng.get('base'):,} – {rng.get('high'):,} (low/base/high)"]
    if fee:
        parts.append(f"indicative fee {fee.get('low'):,}–{fee.get('high'):,}")
    if pb.get("confidence_label"):
        parts.append(pb["confidence_label"])
    return "; ".join(parts)


def _fields(report: dict) -> list[tuple[str, list[tuple[str, Any]]]]:
    prospect = report["prospect"]
    v3 = report.get("v3_analysis") or {}
    snapshot = report.get("snapshot") or {}
    research = report.get("research") or {}
    breakdown = research.get("score_breakdown") or prospect.get("score_breakdown") or {}
    if isinstance(breakdown, str):
        try:
            breakdown = json.loads(breakdown)
        except ValueError:
            breakdown = {}
    pb = prospect.get("prize_breakdown") or {}
    if isinstance(pb, str):
        try:
            pb = json.loads(pb)
        except ValueError:
            pb = {}
    return [
        ("Research provenance", [
            ("Report generated", research.get("generated_at")),
            ("Research as of (last enrichment)", research.get("last_enriched_at")),
            ("Score policy", research.get("score_policy_version") or prospect.get("score_policy_version")),
            ("Latest evidence date", breakdown.get("latest_evidence_date")),
            ("Days since latest evidence", breakdown.get("days_since_latest_evidence")),
            ("Signals used / excluded", f"{breakdown.get('signals_used', '')} / {breakdown.get('signals_excluded', '')}" if breakdown else None),
            ("Estimate assumptions", "\n".join(pb.get("assumptions") or []) or None),
        ]),
        ("Workflow", [
            ("Owner", prospect.get("owner")),
            ("Opportunity stage", prospect.get("opportunity_stage") if prospect.get("opportunity_stage") not in (None, "none") else None),
            ("Next action", prospect.get("next_action")),
            ("Next action due", prospect.get("next_action_due")),
            ("Last contacted", prospect.get("last_contacted_at")),
            ("Outcome reason", prospect.get("outcome_reason")),
        ]),
        ("Company profile", [
            ("Company", prospect.get("company_name")),
            ("Ticker", prospect.get("ticker")),
            ("Sector", prospect.get("gics_sector")),
            ("Industry group", prospect.get("gics_industry_group")),
            ("Website", prospect.get("website")),
            ("Principal activities", prospect.get("principal_activities")),
            ("City", prospect.get("registered_city")),
            ("State", prospect.get("registered_state")),
        ]),
        ("Qualification", [
            ("Prospect score", prospect.get("prospect_score")),
            ("Likelihood score", prospect.get("likelihood_score")),
            ("Lead tier", prospect.get("lead_tier")),
            ("Status", prospect.get("status")),
            ("Estimated impact (AUD, client improvement)", _prize_summary(prospect)),
            ("Strategic direction", prospect.get("strategic_direction")),
            ("Primary tailwind", prospect.get("primary_tailwind")),
            ("Primary headwind", prospect.get("primary_headwind")),
        ]),
        ("Intelligence", [
            ("Key pressures", prospect.get("key_pressures")),
            ("New Delta fit", prospect.get("nd_fit_assessment")),
            ("Outreach hypothesis", prospect.get("outreach_hypothesis")),
            ("Red flags", prospect.get("red_flags")),
            ("Analyst notes", prospect.get("analyst_notes")),
            ("Latest intelligence summary", v3.get("summary")),
            ("Latest intelligence output", v3.get("output_json")),
            ("Commercial snapshot", snapshot.get("snapshot_markdown")),
        ]),
    ]


def report_rows(report: dict) -> list[list[str]]:
    """Create a normalized CSV that preserves every report value."""
    rows = [["Section", "Field", "Value"]]
    for section, fields in _fields(report):
        for label, value in fields:
            rendered = display_value(value)
            if rendered:
                rows.append([section, label, rendered])

    for number, signal in enumerate(report.get("signals") or [], start=1):
        prefix = f"Pressure signal {number}"
        for label, value in [
            ("Pillar", signal.get("pressure_type")),
            ("Strength", signal.get("strength")),
            ("Summary", signal.get("summary")),
            ("Confidence", signal.get("confidence_score")),
            ("Source title", signal.get("source_title")),
            ("Source URL", signal.get("source_url")),
            ("Source date", signal.get("source_date") or "unknown"),
            ("Verification", signal.get("verification_status")),
            ("Extracted quote", signal.get("extracted_quote")),
        ]:
            rendered = display_value(value)
            if rendered:
                rows.append([prefix, label, rendered])

    for number, contact in enumerate(report.get("contacts") or [], start=1):
        prefix = f"Contact {number}"
        for label, value in [
            ("Name", contact.get("full_name")),
            ("Title", contact.get("title")),
            ("Seniority", contact.get("seniority")),
            ("Departments", contact.get("departments")),
            ("Decision rank", contact.get("decision_rank")),
            ("Email", contact.get("email")),
            ("Email status", contact.get("email_status")),
            ("Phone", contact.get("phone")),
            ("LinkedIn", contact.get("linkedin_url")),
            ("Source", contact.get("source_provider")),
            ("Contact type", contact.get("contact_kind")),
            ("Confidence", contact.get("confidence_score")),
            ("Last verified", contact.get("last_verified_at")),
        ]:
            rendered = display_value(value)
            if rendered:
                rows.append([prefix, label, rendered])
    return rows


def csv_safe(value: str) -> str:
    """Prefix values a spreadsheet would execute as a formula."""
    return f"'{value}" if value and value[0] in ("=", "+", "-", "@", "\t", "\r") else value


def build_csv(report: dict) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerows([[csv_safe(cell) for cell in row] for row in report_rows(report)])
    return output.getvalue().encode("utf-8-sig")


def _paragraph_text(value: Any) -> str:
    text = display_value(value)
    # Built-in ReportLab fonts do not reliably cover every Unicode character.
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()


def build_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    prospect = report["prospect"]
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=LETTER, leftMargin=0.65 * inch, rightMargin=0.65 * inch,
        topMargin=0.65 * inch, bottomMargin=0.65 * inch,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.black)
    subtitle = ParagraphStyle("ReportSubtitle", parent=styles["Normal"], alignment=TA_CENTER, fontSize=9, leading=12, textColor=colors.HexColor("#505050"))
    heading = ParagraphStyle("ReportHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=colors.black, spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle("ReportBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, textColor=colors.black)
    small = ParagraphStyle("ReportSmall", parent=body, fontSize=8, leading=10)
    story = [
        Paragraph(escape(_paragraph_text(prospect.get("company_name") or "Company intelligence report")), title),
        Paragraph(escape(f"{prospect.get('ticker') or ''} | Generated from Delta Prospect System"), subtitle),
        Spacer(1, 12),
    ]

    for section, fields in _fields(report):
        populated = [(label, value) for label, value in fields if display_value(value)]
        if not populated:
            continue
        story.append(Paragraph(section, heading))
        table_data = [[Paragraph(f"<b>{escape(label)}</b>", small), Paragraph(escape(_paragraph_text(value)).replace("\n", "<br/>"), small)] for label, value in populated]
        table = Table(table_data, colWidths=[1.45 * inch, 5.2 * inch], repeatRows=0)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9D9D9")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F3F3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(table)

    signals = report.get("signals") or []
    if signals:
        story.append(Paragraph("Pressure signals", heading))
        signal_rows = [[Paragraph("Pillar", small), Paragraph("Strength", small), Paragraph("Summary", small), Paragraph("Source", small), Paragraph("Date", small)]]
        for signal in signals:
            signal_rows.append([
                Paragraph(escape(_paragraph_text(signal.get("pressure_type"))), small),
                Paragraph(escape(_paragraph_text(signal.get("strength"))), small),
                Paragraph(escape(_paragraph_text(signal.get("summary"))).replace("\n", "<br/>"), small),
                Paragraph(escape(_paragraph_text(signal.get("source_title"))) + (f"<br/><i>{escape(_paragraph_text(signal.get('verification_status')))}</i>" if signal.get("verification_status") else ""), small),
                Paragraph(escape(_paragraph_text(signal.get("source_date") or "unknown")), small),
            ])
        table = Table(signal_rows, colWidths=[0.85 * inch, 0.7 * inch, 3.1 * inch, 1.35 * inch, 0.7 * inch], repeatRows=1)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9D9D9")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#303030")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)

    contacts = report.get("contacts") or []
    if contacts:
        story.append(Paragraph("Contacts", heading))
        contact_rows = [[Paragraph("Name", small), Paragraph("Title", small), Paragraph("Email", small), Paragraph("Phone", small), Paragraph("Rank", small)]]
        for contact in contacts:
            contact_rows.append([
                Paragraph(escape(_paragraph_text(contact.get("full_name"))) + (f"<br/><i>{escape(_paragraph_text(contact.get('contact_kind')))}</i>" if contact.get("contact_kind") else ""), small),
                Paragraph(escape(_paragraph_text(contact.get("title"))), small),
                Paragraph(escape(_paragraph_text(contact.get("email"))), small),
                Paragraph(escape(_paragraph_text(contact.get("phone"))), small),
                Paragraph(escape(_paragraph_text(contact.get("decision_rank"))), small),
            ])
        table = Table(contact_rows, colWidths=[1.3 * inch, 2.1 * inch, 1.65 * inch, 1.15 * inch, 0.45 * inch], repeatRows=1)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9D9D9")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#303030")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)

    doc.build(story)
    return output.getvalue()


def build_docx(report: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    def set_cell_shading(cell, fill: str):
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    def set_cell_text(cell, value: Any, *, bold: bool = False, color: str | None = None):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(display_value(value))
        run.bold = bold
        run.font.size = Pt(9)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    prospect = report["prospect"]
    heading = document.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run(prospect.get("company_name") or "Company intelligence report")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"{prospect.get('ticker') or ''} | Generated from Delta Prospect System").italic = True

    for section_name, fields in _fields(report):
        populated = [(label, value) for label, value in fields if display_value(value)]
        if not populated:
            continue
        document.add_heading(section_name, level=1)
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in populated:
            cells = table.add_row().cells
            set_cell_shading(cells[0], "F3F3F3")
            set_cell_text(cells[0], label, bold=True)
            set_cell_text(cells[1], value)

    for section_name, records, columns in [
        ("Pressure signals", report.get("signals") or [], [("Pillar", "pressure_type"), ("Strength", "strength"), ("Summary", "summary"), ("Source", "source_title"), ("Date", "source_date"), ("Verification", "verification_status")]),
        ("Contacts", report.get("contacts") or [], [("Name", "full_name"), ("Title", "title"), ("Email", "email"), ("Phone", "phone"), ("Type", "contact_kind"), ("Rank", "decision_rank")]),
    ]:
        if not records:
            continue
        document.add_heading(section_name, level=1)
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for cell, (label, _) in zip(table.rows[0].cells, columns):
            set_cell_shading(cell, "303030")
            set_cell_text(cell, label, bold=True, color="FFFFFF")
        for record in records:
            for cell, (_, key) in zip(table.add_row().cells, columns):
                set_cell_text(cell, record.get(key))

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
